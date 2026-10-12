#!/usr/bin/env python3
"""Generate supabase/seed.sql from the Garuda workbook.

    python3 -I scripts/import_xlsx.py [--xlsx PATH] [--seed PATH] [--rows PATH] [--check]

Default workbook: "Tracker Tindak Lanjut People Development Garuda - DUMMY.xlsx" in the repo root.

Writes
  supabase/seed.sql      settings, stations (adds missing codes), staff, weekly_scores, bmi_checks,
                         action_items, remedials, replacements. One transaction. Truncates the
                         data tables first (a seed is the baseline, not a merge).
  scripts/out/rows.json  the same rows in camelCase + settings + the workbook's cached Dashboard and
                         Laporan Mingguan results, for the parity test (lib/workbook.test.ts).
Both may contain personal data and are gitignored.

Reads cached values (data_only=True), so the workbook must have been saved by Excel, Google Sheets
or LibreOffice (or baked), not straight from openpyxl.

--check only validates and prints a short report, writing nothing.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
import openpyxl  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "Tracker Tindak Lanjut People Development Garuda - DUMMY.xlsx"
DEFAULT_SEED = ROOT / "supabase" / "seed.sql"
DEFAULT_ROWS = ROOT / "scripts" / "out" / "rows.json"

PROGRAMS = ("IULS", "RUST")
TOUCH = ("Check-in Counter", "Customer Service", "Boarding Gate", "Transfer Desk", "Baggage Service")
ASSIGNMENT = ("Aktif", "Remedial", "Individual Coaching", "Ditarik dari layanan")
GROUPS = ("Coaching Frontliner", "Timeline Training", "Evaluasi IULS & RUST")
ASPECTS = ("Fit & Grooming (BMI)", "Product Knowledge", "Komunikasi & Interactive Skill",
           "Grooming & Aksesori", "Coaching Coverage", "Remedial")
STATUS = ("Belum Mulai", "On Progress", "Selesai", "Tertunda")
KINDS = ("Sekali", "Rutin")
REMEDIAL = ("Refresher & Remedial Assessment", "Individual Coaching", "Refresher + Individual Coaching")
YA_BELUM = ("Ya", "Belum")

# Parameter label (column B) -> settings column. Row positions may move; labels are the key.
PARAMETER_LABELS = {
    "Tanggal mulai Minggu ke-1 (Senin)": ("week1_start", "date"),
    "Jumlah minggu pemantauan": ("weeks", "int"),
    "Tanggal cek BMI slot 1": ("bmi_first_check", "date"),
    "Interval cek BMI – ada catatan (hari)": ("bmi_interval_days", "int"),
    "Interval cek BMI – tanpa catatan (bulan)": ("bmi_interval_months", "int"),
    "Jumlah slot cek di sheet BMI": ("bmi_periods", "int"),
    "BMI kategori Kurus jika <": ("bmi_underweight_below", "num"),
    "BMI Normal s.d.": ("bmi_normal_max", "num"),
    "BMI Overweight s.d. (di atasnya Obesitas)": ("bmi_overweight_max", "num"),
    "Passing grade (skor ≥)": ("pass_grade", "num"),
    "Bobot Role Play di Final Result": ("weight_role_play", "num"),
    "Bobot Post-Test di Final Result": ("weight_post_test", "num"),
    "Target proporsi BMI overweight (maks.)": ("target_overweight_max", "num"),
    "Jangka waktu target BMI (bulan)": ("target_overweight_months", "int"),
    "Target ketepatan jawaban Product Knowledge": ("target_knowledge", "num"),
    "Target skor Interactive Communicating": ("target_interactive", "num"),
    "Target cakupan coaching (jumlah frontliner)": ("target_coverage", "int"),
    "Frontliner sudah coaching (update manual)": ("coverage_done", "int"),
    "Target kelulusan peserta IULS/RUST setelah remedial": ("target_pass_rate", "num"),
    "Batas waktu penggantian SDM": ("replacement_deadline", "date"),
}

BMI_FIRST_COL = 8  # H: Tgl Cek of slot 1; 5 columns per slot (date, height, weight, BMI, category)
BMI_NOTE_COL = 45  # AS: Catatan Program Kebugaran


class ImportFailure(Exception):
    pass


def fail(where: str, msg: str) -> None:
    raise ImportFailure(f"{where}: {msg}")


def text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def choice(v, allowed: tuple[str, ...], where: str) -> str | None:
    s = text(v)
    if s is None:
        return None
    for a in allowed:
        if a.lower() == s.lower():
            return a
    fail(where, f"{s!r} is not one of {allowed}")
    return None


def number(v, where: str, lo: float | None = None, hi: float | None = None):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        fail(where, f"{v!r} is not a number")
    if (lo is not None and n < lo) or (hi is not None and n > hi):
        fail(where, f"{n} outside {lo}..{hi}")
    return int(n) if n.is_integer() else n


def iso_date(v, where: str) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    fail(where, f"{v!r} is not a date")
    return None


def camel(key: str) -> str:
    head, *rest = key.split("_")
    return head + "".join(p[:1].upper() + p[1:] for p in rest)


def sql_literal(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"


def insert(table: str, cols: list[str], rows: list[dict]) -> str:
    if not rows:
        return f"-- {table}: no rows\n"
    values = ",\n".join("  (" + ", ".join(sql_literal(r[c]) for c in cols) + ")" for r in rows)
    return f"insert into public.{table} ({', '.join(cols)}) values\n{values};\n"


# ---------------------------------------------------------------- readers


def read_settings(wb) -> dict:
    ws = wb["Parameter"]
    out: dict = {}
    for r in range(5, ws.max_row + 1):
        label = text(ws.cell(r, 2).value)
        if label in PARAMETER_LABELS:
            col, kind = PARAMETER_LABELS[label]
            v = ws.cell(r, 3).value
            where = f"Parameter!C{r}"
            out[col] = iso_date(v, where) if kind == "date" else number(v, where)
            if kind == "int" and out[col] is not None:
                out[col] = int(out[col])
    missing = [c for c, _ in PARAMETER_LABELS.values() if out.get(c) is None]
    if missing:
        fail("Parameter", f"missing values for {missing}")
    return out


def read_stations(wb) -> list[str]:
    ws = wb["Parameter"]
    head = next(c for c in range(1, ws.max_column + 1) if text(ws.cell(4, c).value) == "Unit / Station")
    out = []
    for r in range(5, ws.max_row + 1):
        code = text(ws.cell(r, head).value)
        if code:
            out.append(code.upper())
    return out


def read_staff(wb) -> list[dict]:
    ws = wb["Master SDM"]
    bm = wb["Cek BMI 2 Mingguan"]
    rows = []
    seen = set()
    for r in range(5, ws.max_row + 1):
        code = text(ws.cell(r, 2).value)
        name = text(ws.cell(r, 6).value)
        if not code and not name:
            continue
        where = f"Master SDM!row {r}"
        if not code or not name:
            fail(where, "ID SDM and Nama must both be filled")
        if code in seen:
            fail(where, f"duplicate ID {code}")
        seen.add(code)
        has_note = text(bm.cell(r + 1, 7).value)
        rows.append({
            "code": code,
            "program": choice(ws.cell(r, 3).value, PROGRAMS, where + " Program"),
            "provider": text(ws.cell(r, 4).value),
            "station": (text(ws.cell(r, 5).value) or fail(where, "Unit is empty") or "").upper(),
            "name": name,
            "employee_no": text(ws.cell(r, 7).value),
            "gender": choice(ws.cell(r, 8).value, ("L", "P"), where + " L/P"),
            "touch_point": choice(ws.cell(r, 9).value, TOUCH, where + " Touch Point"),
            "role_play": number(ws.cell(r, 10).value, where + " Role Play", 0, 100),
            "post_test": number(ws.cell(r, 11).value, where + " Post-Test", 0, 100),
            "report_remark": choice(ws.cell(r, 13).value, ("Passed", "Failed"), where + " Remark"),
            "assignment_status": choice(ws.cell(r, 17).value, ASSIGNMENT, where + " Status") or "Aktif",
            "feedback": text(ws.cell(r, 18).value),
            "notes": text(ws.cell(r, 19).value),
            "bmi_has_note": None if has_note is None else has_note.lower() == "ya",
            "bmi_note": text(bm.cell(r + 1, BMI_NOTE_COL).value),
            "sort_order": len(rows) + 1,
        })
    return rows


def read_weekly_scores(wb, codes: set[str]) -> list[dict]:
    ws = wb["Log Performa Mingguan"]
    rows = {}
    for r in range(5, ws.max_row + 1):
        week, code = ws.cell(r, 1).value, text(ws.cell(r, 3).value)
        vals = [ws.cell(r, c).value for c in (7, 8, 9, 10, 11, 15, 16, 17)]
        if all(v is None or text(v) is None for v in vals):
            continue
        where = f"Log Performa Mingguan!row {r}"
        if week is None or code is None:
            fail(where, "Minggu Ke and ID SDM are required")
        if code not in codes:
            fail(where, f"unknown ID {code}")
        key = (code, int(week))
        if key in rows:
            fail(where, f"duplicate week {week} for {code}")
        rows[key] = {
            "staff_code": code,
            "week": int(week),
            "touch_point": choice(vals[0], TOUCH, where + " Touch Point"),
            "score_grooming": number(vals[1], where + " H", 0, 100),
            "score_knowledge": number(vals[2], where + " I", 0, 100),
            "score_attitude": number(vals[3], where + " J", 0, 100),
            "score_communication": number(vals[4], where + " K", 0, 100),
            "gap_topic": text(vals[5]),
            "observer": text(vals[6]),
            "coaching_notes": text(vals[7]),
        }
    return list(rows.values())


def read_bmi(wb, codes: set[str], slots: int) -> list[dict]:
    ws = wb["Cek BMI 2 Mingguan"]
    out = []
    for r in range(6, ws.max_row + 1):
        code = text(ws.cell(r, 2).value)
        if not code:
            continue
        for k in range(slots):
            c0 = BMI_FIRST_COL + 5 * k
            date, h, w = (ws.cell(r, c0 + i).value for i in range(3))
            if h is None and w is None:
                continue
            where = f"Cek BMI 2 Mingguan!row {r} slot {k + 1}"
            if code not in codes:
                fail(where, f"unknown ID {code}")
            if h is None or w is None:
                fail(where, "Tinggi and Berat must both be filled")
            out.append({
                "staff_code": code,
                "period": k + 1,
                "check_date": iso_date(date, where),
                "height_cm": number(h, where + " tinggi", 120, 210),
                "weight_kg": number(w, where + " berat", 30, 200),
            })
    return out


def read_action_items(wb) -> list[dict]:
    ws = wb["Tindak Lanjut"]
    out = []
    for r in range(5, ws.max_row + 1):
        code = text(ws.cell(r, 2).value)
        if not code:
            continue
        where = f"Tindak Lanjut!row {r}"
        progress = number(ws.cell(r, 13).value, where + " % Progres", 0, 1) or 0
        out.append({
            "code": code,
            "group_name": choice(ws.cell(r, 3).value, GROUPS, where + " Kelompok") or fail(where, "Kelompok empty"),
            "aspect": choice(ws.cell(r, 4).value, ASPECTS, where + " Aspek"),
            "root_cause": text(ws.cell(r, 5).value),
            "action": text(ws.cell(r, 6).value),
            "target": text(ws.cell(r, 7).value),
            "kind": choice(ws.cell(r, 8).value, KINDS, where + " Jenis"),
            "schedule": text(ws.cell(r, 9).value),
            "due_date": iso_date(ws.cell(r, 10).value, where + " Batas Waktu"),
            "pic": text(ws.cell(r, 11).value),
            "status": choice(ws.cell(r, 12).value, STATUS, where + " Status") or "Belum Mulai",
            "progress": int(round(progress * 100)),
            "updated_on": iso_date(ws.cell(r, 14).value, where + " Tgl Update"),
            "evidence": text(ws.cell(r, 15).value),
            "kps_notes": text(ws.cell(r, 16).value),
            "source": text(ws.cell(r, 19).value),
            "sort_order": len(out) + 1,
        })
    return out


def read_remedials(wb, codes: set[str]) -> list[dict]:
    ws = wb["Remedial"]
    out, seen = [], set()
    for r in range(5, ws.max_row + 1):
        code = text(ws.cell(r, 2).value)
        vals = [ws.cell(r, c).value for c in (8, 9, 10, 11, 15, 16)]
        if not code:
            continue
        where = f"Remedial!row {r}"
        if code not in codes:
            fail(where, f"unknown ID {code}")
        if code in seen:
            fail(where, f"duplicate ID {code}")
        seen.add(code)
        out.append({
            "staff_code": code,
            "action": choice(vals[0], REMEDIAL, where + " Tindakan"),
            "remedial_on": iso_date(vals[1], where + " Tgl"),
            "role_play": number(vals[2], where + " Role Play", 0, 100),
            "post_test": number(vals[3], where + " Post-Test", 0, 100),
            "reported": choice(vals[4], YA_BELUM, where + " Dilaporkan"),
            "notes": text(vals[5]),
            "sort_order": len(out) + 1,
        })
    return out


def read_replacements(wb, staff: list[dict]) -> list[dict]:
    ws = wb["Penggantian SDM"]
    by_name = {}
    for s in staff:
        by_name.setdefault((s["station"], s["name"].lower()), s["code"])
    out = []
    for r in range(5, ws.max_row + 1):
        name = text(ws.cell(r, 4).value)
        if not name:
            continue
        where = f"Penggantian SDM!row {r}"
        station = (text(ws.cell(r, 3).value) or "").upper() or None
        out.append({
            "program": choice(ws.cell(r, 2).value, PROGRAMS, where + " Program"),
            "station": station,
            "staff_code": by_name.get((station, name.lower())),
            "replaced_name": name,
            "reason": text(ws.cell(r, 5).value),
            "withdrawn_on": iso_date(ws.cell(r, 6).value, where + " Tgl Ditarik"),
            "replacement_name": text(ws.cell(r, 7).value),
            "effective_on": iso_date(ws.cell(r, 8).value, where + " Tgl Efektif"),
            "training_on": iso_date(ws.cell(r, 9).value, where + " Tgl Training"),
            "post_test": number(ws.cell(r, 10).value, where + " Post-Test", 0, 100),
            "final_result": number(ws.cell(r, 11).value, where + " Final", 0, 100),
            "reported": choice(ws.cell(r, 14).value, YA_BELUM, where + " Dilaporkan"),
            "notes": text(ws.cell(r, 15).value),
            "sort_order": len(out) + 1,
        })
    return out


# ---------------------------------------------------------------- cached results (parity test)


def find_row(ws, col: int, label: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if text(ws.cell(r, col).value) == label:
            return r
    fail(ws.title, f"label {label!r} not found in column {col}")
    return 0


def section_rows(ws, title_prefix: str, keys: list[str]) -> dict[str, list]:
    """Rows under the section whose title in column B starts with title_prefix, keyed by column B."""
    top = next(r for r in range(1, ws.max_row + 1) if (text(ws.cell(r, 2).value) or "").startswith(title_prefix))
    out = {}
    for key in keys:
        r = find_row(ws, 2, key, top + 2)
        out[key] = [ws.cell(r, c).value for c in range(3, 16)]
    return out


def expected(wb, stations: list[str]) -> dict:
    d, lm = wb["Dashboard"], wb["Laporan Mingguan"]
    kpi = {}
    for r in range(7, 14):
        kpi[text(d.cell(r, 2).value)] = [d.cell(r, c).value for c in (3, 4, 5)]
    units = [text(d.cell(r, 2).value) for r in range(1, d.max_row + 1)]
    unit_keys = [u for u in stations if u in units] + ["TOTAL"]
    return {
        "today": iso_date(d["C3"].value, "Dashboard!C3"),
        "week": int(lm["C4"].value),
        "dashboard": {
            "kpis": kpi,
            "byGroup": section_rows(d, "2. PROGRES", list(GROUPS) + ["TOTAL"]),
            "byAspect": section_rows(d, "3. PROGRES", list(ASPECTS) + ["TOTAL"]),
            "programs": section_rows(d, "4. KELULUSAN", ["IULS", "RUST", "TOTAL"]),
            "units": section_rows(d, "5. PERFORMA", unit_keys),
        },
        "laporan": {
            "performance": section_rows(lm, "A. PERFORMA", ["IULS", "RUST", "TOTAL"]),
            "aspects": section_rows(lm, "A2.", ["IULS", "RUST", "TOTAL"]),
            "actions": section_rows(lm, "B. PROGRES", list(GROUPS) + ["TOTAL"]),
            "bmi": section_rows(lm, "C. CEK BMI", ["IULS", "RUST", "TOTAL"]),
            "trend": section_rows(lm, "D. TREN", ["IULS", "RUST", "SEMUA"]),
        },
    }


# ---------------------------------------------------------------- main

COLS = {
    "staff": ["code", "program", "provider", "station", "name", "employee_no", "gender", "touch_point",
              "role_play", "post_test", "report_remark", "assignment_status", "feedback", "notes",
              "bmi_has_note", "bmi_note", "sort_order"],
    "weekly_scores": ["staff_code", "week", "touch_point", "score_grooming", "score_knowledge", "score_attitude",
                      "score_communication", "gap_topic", "observer", "coaching_notes"],
    "bmi_checks": ["staff_code", "period", "check_date", "height_cm", "weight_kg"],
    "action_items": ["code", "group_name", "aspect", "root_cause", "action", "target", "kind", "schedule",
                     "due_date", "pic", "status", "progress", "updated_on", "evidence", "kps_notes", "source",
                     "sort_order"],
    "remedials": ["staff_code", "action", "remedial_on", "role_play", "post_test", "reported", "notes", "sort_order"],
    "replacements": ["program", "station", "staff_code", "replaced_name", "reason", "withdrawn_on",
                     "replacement_name", "effective_on", "training_on", "post_test", "final_result", "reported",
                     "notes", "sort_order"],
}


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--seed", type=Path, default=DEFAULT_SEED)
    ap.add_argument("--rows", type=Path, default=DEFAULT_ROWS)
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args(argv)

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    if wb["Master SDM"]["L5"].value is None and wb["Master SDM"]["B5"].value is not None:
        raise ImportFailure("the workbook has no cached formula values; open and save it in Excel/Sheets first")

    settings = read_settings(wb)
    stations = read_stations(wb)
    staff = read_staff(wb)
    codes = {s["code"] for s in staff}
    for s in staff:
        if s["station"] not in stations:
            stations.append(s["station"])
    data = {
        "staff": staff,
        "weekly_scores": read_weekly_scores(wb, codes),
        "bmi_checks": read_bmi(wb, codes, settings["bmi_periods"]),
        "action_items": read_action_items(wb),
        "remedials": read_remedials(wb, codes),
        "replacements": read_replacements(wb, staff),
    }

    print(f"workbook: {a.xlsx.name}")
    for k, v in data.items():
        print(f"  {k:<15} {len(v)}")
    dummy = sum(1 for s in staff if "DUMMY" in (s["notes"] or "").upper())
    if dummy:
        print(f"  note: {dummy} staff rows are marked DUMMY")
    if a.check:
        return 0

    seed = [
        f"-- Generated by scripts/import_xlsx.py from {a.xlsx.name}. Do not edit; do not commit.",
        "begin;",
        "truncate public.weekly_scores, public.bmi_checks, public.remedials, public.replacements,"
        " public.weekly_reports, public.action_items, public.staff restart identity cascade;",
        "insert into public.stations (code, sort) values\n"
        + ",\n".join(f"  ({sql_literal(c)}, {i + 1})" for i, c in enumerate(stations))
        + "\non conflict (code) do update set sort = excluded.sort;",
        "update public.settings set "
        + ", ".join(f"{k} = {sql_literal(v)}" for k, v in settings.items())
        + " where id = 1;",
    ]
    for table in COLS:
        seed.append(insert(table, COLS[table], data[table]))
    seed.append("commit;")
    a.seed.parent.mkdir(parents=True, exist_ok=True)
    a.seed.write_text("\n".join(seed) + "\n", encoding="utf-8")

    rows = {
        "settings": {camel(k): v for k, v in settings.items()},
        "stations": [{"code": c, "sort": i + 1} for i, c in enumerate(stations)],
        **{camel(t): [{camel(k): v for k, v in r.items()} for r in rs] for t, rs in data.items()},
        "expected": expected(wb, stations),
    }
    a.rows.parent.mkdir(parents=True, exist_ok=True)
    (a.rows.parent / ".gitignore").write_text("*\n", encoding="utf-8")
    a.rows.write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(f"wrote {a.seed.relative_to(ROOT)} and {a.rows.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except ImportFailure as e:
        print(f"error: {e}", file=sys.stderr)
        sys.exit(1)
